from __future__ import annotations

import argparse
import json
from datetime import datetime, timezone
from pathlib import Path

from common import (
    display_value,
    ensure_within,
    inspect_template,
    load_facts,
    load_json,
    matching_configs,
    next_config_version,
    resolved,
    target_cell,
    validate_config,
    write_json_new,
)
from docx import Document
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


FIELD_DEFINITIONS = [
    ("application_date", "申请日期", "申请日期", 0, 1, {"type": "direct", "mode": "cell"}),
    ("acceptance_number", "项目受理号", "项目受理号", 0, 5, {"type": "direct", "mode": "cell"}),
    ("protocol_number", "方案号", "方案编号", 1, 1, {"type": "direct", "mode": "cell"}),
    ("protocol_version", "方案版本号", "方案版本号", 1, 5, {"type": "direct", "mode": "cell"}),
    ("protocol_title", "方案名称", "方案名称", 2, 1, {"type": "direct", "mode": "cell"}),
    ("principal_investigator", "主要研究者", "主要研究者", 3, 1, {"type": "direct", "mode": "cell"}),
    ("department", "研究科室", "研究科室", 4, 1, {"type": "direct", "mode": "cell"}),
    ("external_investigators", "外单位协作研究者", "外单位协作研究者", 5, 1, {"type": "direct", "mode": "cell"}),
    ("sponsor", "申办者", "申办者", 6, 1, {"type": "direct", "mode": "cell"}),
    ("service_provider", "第三方服务公司", "第三方服务公司", 7, 1, {"type": "direct", "mode": "cell"}),
    ("enrollment", "本中心计划入组人数／研究参与者总人数", "本中心计划入组人数／受试者总人数", 8, 3, {"type": "direct", "mode": "cell"}),
    ("site_count", "研究中心数目", "研究中心数目", 9, 2, {"type": "direct", "mode": "cell"}),
    ("study_period", "研究起止期限（年/月）", "研究起止期限（年/月）", 10, 2, {"type": "direct", "mode": "cell"}),
    ("funding_source", "经费来源", "经费来源", 11, 2, {"type": "choice", "unchecked": "口", "checked": "☒", "value_to_option": {"政府立项": "政府立项", "学会/协会": "学会/协会", "医药公司": "医药公司"}}),
    ("prior_ethics_rejection", "其他伦理委员会拒绝或否决", "是否曾递交其他伦理委员会并被拒绝或否决", 12, 2, {"type": "choice", "unchecked": "口", "checked": "☒", "value_to_option": {"是": "是", "否": "否"}}),
    ("human_genetics_outbound", "生物样本或人遗数据出本单位", "是否有生物样本或涉及人类遗传资源数据出本单位", 13, 0, {"type": "choice", "unchecked": "口", "checked": "☒", "value_to_option": {"是": "是", "否": "否"}}),
    ("special_population_collection", "特定人群样本和信息采集", "是否涉及家系或特定地区人群的生物样本和信息采集", 14, 0, {"type": "choice", "unchecked": "口", "checked": "☒", "value_to_option": {"是": "是", "否": "否"}}),
    ("study_type", "研究类型", "研究类型", 15, 0, {"type": "choice", "unchecked": "口", "checked": "☒", "value_to_option": {"药物注册临床试验": "药物注册临床试验", "医疗器械注册临床试验": "医疗器械注册临床试验", "生物样本/数据库": "生物样本/数据库"}}),
    ("publication_method", "临床研究结果报告和发表的方式", "临床研究结果报告和发表的方式", 16, 0, {"type": "direct", "mode": "inline_after_anchor", "anchor": "临床研究结果报告和发表的方式："}),
    ("pi_signature", "主要研究者签名及日期", None, 17, 0, {"type": "manual"}),
    ("department_head_signature", "科室主任签名及日期", None, 18, 0, {"type": "manual"}),
    ("ethics_committee_area", "伦理委员会填写区域", None, 19, 0, {"type": "manual"}),
]


def draft_xlsx_mapping(target: Path, facts_path: Path, config_library: Path, inspection: dict, mapping_spec_path: Path) -> dict:
    mapping_spec = load_json(mapping_spec_path)
    config_id = str(mapping_spec.get("config_id", "")).strip()
    form_type = str(mapping_spec.get("form_type", "")).strip()
    definitions = mapping_spec.get("fields")
    if not config_id or not form_type or not isinstance(definitions, list) or not definitions:
        raise ValueError("Excel 映射说明必须包含 config_id、form_type 和非空 fields")
    facts = load_facts(facts_path)
    workbook = load_workbook(target, data_only=False, read_only=False, keep_links=True)
    fields = []
    review = []
    try:
        for definition in definitions:
            if not isinstance(definition, dict):
                raise ValueError("Excel 映射字段必须为对象")
            target_spec = definition.get("target", {})
            try:
                sheet = workbook[target_spec["sheet"]]
                cell = sheet[target_spec["cell"]]
            except (KeyError, TypeError, ValueError) as exc:
                raise ValueError(f"Excel 映射目标位置无效: {target_spec}") from exc
            if isinstance(cell, MergedCell):
                raise ValueError(f"Excel 映射必须指向合并区域左上角单元格: {target_spec}")
            mapping = definition.get("mapping")
            if not isinstance(mapping, dict) or mapping.get("type") not in {"direct", "manual"}:
                raise ValueError(f"Excel 首版仅允许 direct 或 manual 映射: {definition.get('id', '')}")
            if mapping.get("type") == "direct" and mapping.get("mode") != "cell":
                raise ValueError(f"Excel 直接字段必须使用 cell 模式: {definition.get('id', '')}")
            if mapping.get("type") == "direct" and cell.data_type == "f":
                raise ValueError(f"Excel 映射不得指向公式单元格: {sheet.title}!{cell.coordinate}")
            fact_key = definition.get("fact_key")
            if mapping.get("type") == "direct" and not isinstance(fact_key, str):
                raise ValueError(f"Excel 直接字段缺少 fact_key: {definition.get('id', '')}")
            if mapping.get("type") == "manual":
                fact_key = None
            source = facts.get(fact_key) if fact_key else None
            field = {
                "id": definition.get("id"),
                "label": definition.get("label"),
                "fact_key": fact_key,
                "target": {
                    "sheet": sheet.title,
                    "cell": cell.coordinate,
                    "expected_original": display_value(cell.value),
                },
                "mapping": mapping,
                "mapping_review": "proposed",
                "source_hint": None if source is None else {"source_file": source["source_file"], "source_location": source["source_location"]},
            }
            fields.append(field)
            review.append({
                "id": field["id"],
                "label": field["label"],
                "fact_key": fact_key,
                "target": f"{sheet.title}!{cell.coordinate}",
                "mapping_type": mapping["type"],
                "target_original": display_value(cell.value),
                "source_value": None if source is None else source["value"],
                "source_status": None if source is None else source["status"],
            })
    finally:
        workbook.close()

    version = next_config_version(config_library, config_id)
    payload = {
        "schema_version": "1.0",
        "config_id": config_id,
        "version": version,
        "status": "draft",
        "form_type": form_type,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "template": {key: inspection[key] for key in ("file_name", "sha256", "structure_fingerprint", "format")},
        "fields": fields,
        "approval": {"approved_by": None, "approved_at": None, "approval_id": None},
    }
    validate_config(payload, "draft")
    draft_path = config_library / "drafts" / f"{config_id}-v{version}.json"
    write_json_new(draft_path, payload)
    return {"status": "draft_created", "draft": str(draft_path), "field_count": len(fields), "review": review, "next_action": "停止。向 CRA 展示映射草稿并等待明确批准。"}


def main() -> None:
    parser = argparse.ArgumentParser(description="为首次出现的 CRA Word 或 Excel 模板生成映射配置草稿")
    parser.add_argument("--target", required=True)
    parser.add_argument("--facts", required=True)
    parser.add_argument("--authorized-workspace", required=True)
    parser.add_argument("--config-library", required=True)
    parser.add_argument("--mapping-spec", help="Excel 模板的待审字段位置说明 JSON")
    args = parser.parse_args()

    workspace = resolved(args.authorized_workspace)
    target = resolved(args.target)
    facts_path = resolved(args.facts)
    config_library = resolved(args.config_library)
    paths = [target, facts_path, config_library]
    mapping_spec_path = resolved(args.mapping_spec) if args.mapping_spec else None
    if mapping_spec_path:
        paths.append(mapping_spec_path)
    ensure_within(workspace, paths)

    inspection = inspect_template(target)
    enabled = matching_configs(config_library, inspection["structure_fingerprint"], inspection["sha256"], "enabled")
    if enabled:
        raise ValueError(f"已存在匹配的启用配置，不应生成新草稿: {enabled[0]}")
    existing = matching_configs(config_library, inspection["structure_fingerprint"], inspection["sha256"], "draft")
    if existing:
        print(json.dumps({"status": "existing_draft", "draft": str(existing[0])}, ensure_ascii=False, indent=2))
        return

    if inspection["format"] == "xlsx":
        if mapping_spec_path is None:
            raise ValueError("Excel 新模板必须提供 --mapping-spec，供 CRA 审核具体字段位置")
        print(json.dumps(draft_xlsx_mapping(target, facts_path, config_library, inspection, mapping_spec_path), ensure_ascii=False, indent=2))
        return

    if inspection["table_count"] != 1 or inspection["structure"]["tables"][0]["rows"] != 20:
        raise ValueError("当前 Word 原型只支持已分析的 20 行伦理申请表")

    facts = load_facts(facts_path)
    document = Document(target)
    fields = []
    review = []
    for field_id, label, fact_key, row, column, mapping in FIELD_DEFINITIONS:
        cell = target_cell(document, {"table": 0, "row": row, "column": column})
        source = facts.get(fact_key) if fact_key else None
        field = {
            "id": field_id,
            "label": label,
            "fact_key": fact_key,
            "target": {"table": 0, "row": row, "column": column, "expected_original": cell.text},
            "mapping": mapping,
            "mapping_review": "proposed",
            "source_hint": None if source is None else {"source_file": source["source_file"], "source_location": source["source_location"]},
        }
        fields.append(field)
        review.append({"id": field_id, "label": label, "fact_key": fact_key, "target": f"表1/行{row + 1}/列{column + 1}", "mapping_type": mapping["type"], "source_value": None if source is None else source["value"], "source_status": None if source is None else source["status"]})

    config_id = "xx-hospital-ethics-review-application"
    version = next_config_version(config_library, config_id)
    payload = {
        "schema_version": "1.0",
        "config_id": config_id,
        "version": version,
        "status": "draft",
        "form_type": "ethics-review-application",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "template": {key: inspection[key] for key in ("file_name", "sha256", "structure_fingerprint", "format")},
        "fields": fields,
        "approval": {"approved_by": None, "approved_at": None, "approval_id": None},
    }
    draft_path = config_library / "drafts" / f"{config_id}-v{version}.json"
    write_json_new(draft_path, payload)
    print(json.dumps({"status": "draft_created", "draft": str(draft_path), "field_count": len(fields), "review": review, "next_action": "停止。向 CRA 展示映射草稿并等待明确批准。"}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
